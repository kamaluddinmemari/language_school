import random
from collections import defaultdict
from datetime import datetime, timedelta

import jdatetime
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from rest_framework import status
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.models import User
from accounts.menu_permissions import can_edit_menu
from .attendance import attendance_session_dates_for_slot, jalali_date, session_date_and_sub
from .models import (
    ClassSlot,
    ClassSlotEnrollment,
    ClassSurveyList,
    ClassSurveyQuestion,
    ClassSurveyResponse,
    FRIDAY_EVENT_SLOTS,
    FRIDAY_EVENING_EVENT_SLOTS,
    FRIDAY_MORNING_EVENT_SLOTS,
    THURSDAY_EVENING_EVENT_SLOTS,
    THURSDAY_MORNING_EVENT_SLOTS,
    Term,
    infer_age_group_from_level,
)


DEFAULT_QUESTIONS = [
    'از روند کلی کلاس راضی هستید؟',
    'از نحوه تدریس و برخورد استاد راضی هستید؟',
    'آیا استاد گروه دارد و تکالیف را به‌موقع در گروه قرار می‌دهد؟',
    'آیا کتاب داستان برای نوجوان و آکسفورد و ویدئو برای بزرگسال کار می‌شود؟',
    'از اطلاع‌رسانی‌های مؤسسه رضایت دارید؟',
    'اگر نظر یا انتقادی دارید، بنویسید.',
]

STAFF_ROLES = ('admin', 'office', 'employee')
DEADLINE_MODES = ('none', 'date', 'automatic', 'session')
CALL_STATUSES = ('answered', 'no_answer')


def _age_group(level):
    """ردهٔ سنی سطح را از جدول استاندارد می‌خواند و برای داده‌های قدیمی fallback دارد."""
    value = infer_age_group_from_level(level)
    if value in ('kids', 'teen', 'adult'):
        return value
    code = str(level or '').strip().lower().replace(' ', '')
    if code.startswith(('teen', 'preteen')):
        return 'teen'
    if code[:1] in ('e', 's', 'g', 'i', 'u', 'm', 'h'):
        return 'kids'
    if code.isdigit() or (code and code[0].isdigit()):
        return 'adult'
    return 'all'


def _allowed(request):
    # همان سیاست قبلی: کارشناس آموزش فقط با edit permission و نقش‌های اداریِ اجرایی مجازند.
    return can_edit_menu(request.user, 'class-management') or request.user.role in STAFF_ROLES


def _staff():
    return list(User.objects.filter(role__in=STAFF_ROLES, is_active=True).order_by('first_name', 'last_name', 'id'))


def _staff_payload():
    return [{'id': user.id, 'name': _name(user)} for user in _staff()]


def _name(user):
    return (f'{user.first_name} {user.last_name}'.strip() or user.username) if user else ''


def _ensure_questions():
    if not ClassSurveyQuestion.objects.exists():
        ClassSurveyQuestion.objects.bulk_create([
            ClassSurveyQuestion(age_group='all', text=text, sort_order=index)
            for index, text in enumerate(DEFAULT_QUESTIONS, 1)
        ])


def _questions(age_group=None):
    _ensure_questions()
    qs = ClassSurveyQuestion.objects.filter(active=True)
    if age_group:
        qs = qs.filter(Q(age_group='all') | Q(age_group=age_group))
    return [
        {'id': question.id, 'age_group': question.age_group, 'text': question.text, 'sort_order': question.sort_order}
        for question in qs
    ]


def _jalali_datetime(value):
    if not value:
        return None
    return jdatetime.datetime.fromgregorian(datetime=timezone.localtime(value)).strftime('%Y/%m/%d - %H:%M')


def _iso(value):
    return value.isoformat() if value else None


def _as_bool(value, default=False):
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, int) and value in (0, 1):
        return bool(value)
    if isinstance(value, str):
        normalized = value.strip().lower()
        if normalized in ('true', '1', 'yes'):
            return True
        if normalized in ('false', '0', 'no', ''):
            return False
    raise ValueError('مقدار بولی نامعتبر است')


def _positive_int(value, field_name):
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        raise ValueError(f'{field_name} نامعتبر است')
    if parsed < 1:
        raise ValueError(f'{field_name} نامعتبر است')
    return parsed


def _assigned_filter_value(request):
    raw = request.query_params.get('assigned_to')
    if raw in (None, ''):
        return None, None
    try:
        employee_id = _positive_int(raw, 'شناسه کارمند')
    except ValueError as exc:
        return None, str(exc)
    return employee_id, None


def _row(enrollment, row_number=None):
    slot = enrollment.class_slot
    student = enrollment.student
    result = {
        'key': f'{student.id}:{slot.id}', 'student_id': student.id, 'class_slot_id': slot.id,
        'first_name': student.first_name, 'last_name': student.last_name, 'phone': student.phone or '',
        'gender': student.gender or '', 'level': slot.assigned_level or student.language_level or '',
        'teacher': slot.teacher_name or '', 'day': slot.get_day_type_display(), 'time_slot': slot.time_slot or '',
        'class_number': slot.number, 'age_group': _age_group(slot.assigned_level or student.language_level or ''),
    }
    if row_number is not None:
        result['row_number'] = row_number
    return result


def _survey_slots(survey, responses=None):
    """تمام کلاس‌های survey (نه فقط rows فیلترشده بر اساس کارمند) را یکتا و مرتب برمی‌گرداند."""
    if responses is None:
        responses = survey.responses.select_related('class_slot__term').all()
    slots_by_id = {}
    for response in responses:
        slots_by_id[response.class_slot_id] = response.class_slot
    return [slots_by_id[key] for key in sorted(slots_by_id)]


def _event_time_ranges(slot):
    """بازه‌های استاندارد واقعی جلسات تک‌روزه؛ منبع همان جدول بازه‌های استاندارد مدل است."""
    if slot.day_type == ClassSlot.DayType.THURSDAY_MORNING:
        return THURSDAY_MORNING_EVENT_SLOTS
    if slot.day_type == ClassSlot.DayType.THURSDAY_EVENING:
        return THURSDAY_EVENING_EVENT_SLOTS
    if slot.day_type == ClassSlot.DayType.FRIDAY:
        return FRIDAY_EVENT_SLOTS
    if slot.day_type == ClassSlot.DayType.FRIDAY_MORNING:
        return FRIDAY_MORNING_EVENT_SLOTS
    if slot.day_type == ClassSlot.DayType.FRIDAY_EVENING:
        return FRIDAY_EVENING_EVENT_SLOTS
    return []


def _session_end_time(slot, session_in_day=1):
    ranges = _event_time_ranges(slot)
    # پنجشنبه‌ها سه sub-session مجزا دارند. جمعه در جدول استاندارد فعلی یک بازهٔ معتبر دارد
    # و هر سه رکورد attendance همان زمان پایان واقعی کلاس را دارند.
    if ranges:
        chosen = ranges[min(max(session_in_day, 1), len(ranges)) - 1]
    else:
        chosen = slot.time_slot or ''
    try:
        return datetime.strptime(chosen.split('-', 1)[1].strip(), '%H:%M').time()
    except (IndexError, ValueError):
        return None


def _session_due_at(slot, session_number):
    date_value, session_in_day = session_date_and_sub(slot, session_number, session_count=15)
    if not date_value:
        return None
    end_time = _session_end_time(slot, session_in_day)
    if not end_time:
        return None
    return timezone.make_aware(datetime.combine(date_value, end_time), timezone.get_current_timezone())


def _configured_deadline_at(survey):
    if not survey.deadline_enabled or survey.deadline_mode == 'none':
        return None
    if survey.deadline_mode == 'automatic':
        if not survey.created_at or not survey.deadline_days:
            return None
        return survey.created_at + timedelta(days=survey.deadline_days)
    if survey.deadline_mode == 'date':
        return survey.deadline_at
    return None


def _effective_due_at(response):
    survey = response.survey_list
    if not survey.deadline_enabled or survey.deadline_mode == 'none':
        return None
    if survey.deadline_mode == 'session':
        # زنده محاسبه می‌شود تا تغییر تعطیلات یا تقویم ترم، نمایش calendar/alerts را کهنه نکند.
        return _session_due_at(response.class_slot, survey.deadline_session_number or 0)
    return _configured_deadline_at(survey)


def _session_options(survey, slots=None):
    slots = _survey_slots(survey) if slots is None else slots
    if not slots:
        return []
    lengths = [len(attendance_session_dates_for_slot(slot, session_count=15)) for slot in slots]
    upper_bound = min(lengths) if lengths else 0
    return [
        {'session_number': number, 'label': f'جلسه {number}'}
        for number in range(1, min(15, upper_bound) + 1)
    ]


def _missing_session_slots(survey, session_number, slots=None):
    slots = _survey_slots(survey) if slots is None else slots
    return [
        slot for slot in slots
        if len(attendance_session_dates_for_slot(slot, session_count=15)) < session_number
    ]


def _session_dates_payload(slot):
    dates = attendance_session_dates_for_slot(slot, session_count=15)
    output = []
    seen_by_date = {}
    for index, date_value in enumerate(dates, start=1):
        seen_by_date[date_value] = seen_by_date.get(date_value, 0) + 1
        session_in_day = seen_by_date[date_value]
        end_time = _session_end_time(slot, session_in_day)
        output.append({
            'session_number': index,
            'date': date_value.isoformat(),
            'date_jalali': jalali_date(date_value),
            'session_in_day': session_in_day,
            'end_time': end_time.strftime('%H:%M') if end_time else None,
        })
    return output


def _deadline_payload(survey):
    deadline_at = _configured_deadline_at(survey)
    enabled = bool(survey.deadline_enabled and survey.deadline_mode != 'none')
    return {
        'deadline_enabled': enabled,
        'deadline_mode': survey.deadline_mode if enabled else 'none',
        'deadline_at': _iso(deadline_at),
        'deadline_at_jalali': _jalali_datetime(deadline_at),
        'deadline_days': survey.deadline_days if survey.deadline_mode == 'automatic' and enabled else None,
        'deadline_session_number': survey.deadline_session_number if survey.deadline_mode == 'session' and enabled else None,
    }


def _stored_row_and_number(response):
    key = f'{response.student_id}:{response.class_slot_id}'
    for index, item in enumerate(response.survey_list.rows or [], start=1):
        if item.get('key') == key:
            return item, item.get('row_number') or index
    return {}, None


def _response_payload(response, assigned_to=None, now=None, include_session_dates=False):
    assignee = response.assigned_to
    if assigned_to is not None and (not assignee or assignee.id != assigned_to):
        return None
    item, row_number = _stored_row_and_number(response)
    due_at = _effective_due_at(response)
    now = now or timezone.now()
    payload = {
        **item,
        'key': item.get('key') or f'{response.student_id}:{response.class_slot_id}',
        'student_id': response.student_id,
        'class_slot_id': response.class_slot_id,
        'row_number': row_number,
        'assigned_to': assignee.id if assignee else None,
        'assigned_to_name': _name(assignee),
        'answers': response.answers or {},
        'completed': response.completed,
        'respondent_name': response.respondent_name or '',
        'call_status': response.call_status or '',
        'contacted_at': _iso(response.contacted_at),
        'contacted_at_jalali': _jalali_datetime(response.contacted_at),
        'due_at': _iso(due_at),
        'due_at_jalali': _jalali_datetime(due_at),
        'overdue': bool(due_at and not response.completed and due_at < now),
    }
    if include_session_dates:
        payload['session_dates'] = _session_dates_payload(response.class_slot)
    return payload


def _answer_present(value):
    if isinstance(value, dict):
        for key in ('answer', 'value', 'text'):
            if key in value:
                return _answer_present(value[key])
        return any(_answer_present(item) for item in value.values())
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, (list, tuple)):
        return bool(value)
    return True


def _response_answers_complete(response):
    _ensure_questions()
    age_group = _age_group(response.class_slot.assigned_level or response.student.language_level or '')
    questions = ClassSurveyQuestion.objects.filter(active=True).filter(Q(age_group='all') | Q(age_group=age_group))
    answers = response.answers or {}
    for question in questions:
        value = answers.get(str(question.id), answers.get(question.id))
        if not _answer_present(value):
            return False
    return True


def _responses_for_survey(survey, assigned_to=None):
    queryset = survey.responses.select_related('student', 'class_slot__term', 'assigned_to', 'survey_list__term')
    if assigned_to is not None:
        queryset = queryset.filter(assigned_to_id=assigned_to)
    return list(queryset)


def _summary(survey, responses=None, now=None):
    now = now or timezone.now()
    responses = _responses_for_survey(survey) if responses is None else list(responses)
    overdue_count = sum(1 for response in responses if (due := _effective_due_at(response)) and not response.completed and due < now)
    return {
        'id': survey.id,
        'name': survey.name,
        'term_id': survey.term_id,
        'term_title': survey.term.title,
        'percentage': survey.percentage,
        'use_teacher_class_sampling': survey.use_teacher_class_sampling,
        'teacher_class_percentage': survey.teacher_class_percentage,
        'count': len(responses),
        'completed_count': sum(1 for response in responses if response.completed),
        'assigned_count': sum(1 for response in responses if response.assigned_to_id),
        'overdue_count': overdue_count,
        'created_at': _iso(survey.created_at),
        **_deadline_payload(survey),
        # گزینه‌ها عمداً از همه کلاس‌های خود survey ساخته می‌شوند، نه rows فیلترشدهٔ کارمند.
        'session_options': _session_options(survey),
    }


def _list_payload(request, only_deadline_lists=False):
    assigned_to, error = _assigned_filter_value(request)
    if error:
        return None, error
    now = timezone.now()
    surveys = ClassSurveyList.objects.select_related('term').all()
    term_id = request.query_params.get('term_id')
    if term_id:
        surveys = surveys.filter(term_id=term_id)
    surveys = surveys[:200]
    lists = []
    for survey in surveys:
        if only_deadline_lists and not (survey.deadline_enabled and survey.deadline_mode != 'none'):
            continue
        lists.append(_summary(survey, _responses_for_survey(survey, assigned_to), now))
    return {'lists': lists, 'staff': _staff_payload(), 'server_now': _iso(now)}, None


def _eligible_staff_by_ids(raw_ids):
    if not isinstance(raw_ids, list) or not raw_ids:
        raise ValueError('employee_ids باید یک آرایهٔ غیرخالی باشد')
    ids = []
    for raw_id in raw_ids:
        parsed = _positive_int(raw_id, 'شناسه کارمند')
        if parsed not in ids:
            ids.append(parsed)
    users_by_id = {
        user.id: user
        for user in User.objects.filter(id__in=ids, role__in=STAFF_ROLES, is_active=True)
    }
    missing = [user_id for user_id in ids if user_id not in users_by_id]
    if missing:
        raise ValueError('کارمند انتخاب‌شده فعال نیست یا نقش مجاز ندارد')
    return ids, users_by_id


def _parse_deadline_datetime(value):
    if not isinstance(value, str) or not value.strip():
        raise ValueError('deadline_at با قالب ISO الزامی است')
    parsed = parse_datetime(value.strip())
    if not parsed:
        raise ValueError('deadline_at باید datetime با قالب ISO باشد')
    if timezone.is_naive(parsed):
        parsed = timezone.make_aware(parsed, timezone.get_current_timezone())
    return parsed


def _save_due_dates(survey):
    responses = list(survey.responses.select_related('class_slot__term', 'survey_list__term').all())
    for response in responses:
        response.due_at = _effective_due_at(response)
    if responses:
        ClassSurveyResponse.objects.bulk_update(responses, ['due_at'])
    return responses


class SurveyListView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not _allowed(request):
            return Response({'error': 'دسترسی ندارید'}, status=403)
        list_id = request.query_params.get('id')
        assigned_to, error = _assigned_filter_value(request)
        if error:
            return Response({'error': error}, status=400)
        if list_id:
            try:
                survey = ClassSurveyList.objects.select_related('term').get(pk=int(list_id))
            except (ClassSurveyList.DoesNotExist, ValueError, TypeError):
                return Response({'error': 'لیست پیدا نشد'}, status=404)
            now = timezone.now()
            responses = _responses_for_survey(survey, assigned_to)
            rows = [
                _response_payload(response, now=now, include_session_dates=True)
                for response in responses
            ]
            return Response({
                **_summary(survey, responses, now),
                'rows': rows,
                'questions': _questions(),
                'staff': _staff_payload(),
            })
        payload, error = _list_payload(request)
        if error:
            return Response({'error': error}, status=400)
        payload['questions'] = _questions()
        return Response(payload)

    def post(self, request):
        if not can_edit_menu(request.user, 'class-management'):
            return Response({'error': 'دسترسی ویرایش ندارید'}, status=403)
        try:
            term = Term.objects.get(pk=int(request.data.get('term_id')))
        except (Term.DoesNotExist, ValueError, TypeError):
            return Response({'error': 'ترم پیدا نشد'}, status=400)
        try:
            percentage = int(request.data.get('percentage', 20))
        except (ValueError, TypeError):
            percentage = 20
        if percentage < 1 or percentage > 100:
            return Response({'error': 'درصد باید بین ۱ تا ۱۰۰ باشد'}, status=400)
        try:
            use_teacher_class_sampling = _as_bool(
                request.data.get('use_teacher_class_sampling'), default=False,
            )
        except ValueError as exc:
            return Response({'error': str(exc)}, status=400)
        try:
            teacher_class_percentage = int(request.data.get('teacher_class_percentage', percentage))
        except (ValueError, TypeError):
            teacher_class_percentage = 0
        if use_teacher_class_sampling and not 1 <= teacher_class_percentage <= 100:
            return Response({'error': 'teacher_class_percentage باید بین ۱ تا ۱۰۰ باشد'}, status=400)
        enrollments = list(
            ClassSlotEnrollment.objects.filter(class_slot__term=term, student__role='student')
            .select_related('student', 'class_slot')
        )
        rng = random.SystemRandom()
        if use_teacher_class_sampling:
            enrollments_by_slot = defaultdict(list)
            slots_by_teacher = defaultdict(dict)
            for enrollment in enrollments:
                enrollments_by_slot[enrollment.class_slot_id].append(enrollment)
                teacher_key = (enrollment.class_slot.teacher_name or '').strip() or 'بدون استاد'
                slots_by_teacher[teacher_key][enrollment.class_slot_id] = enrollment.class_slot
            selected_slot_ids = set()
            for slots_by_id in slots_by_teacher.values():
                teacher_slots = list(slots_by_id.values())
                count = max(1, round(len(teacher_slots) * teacher_class_percentage / 100))
                selected_slot_ids.update(slot.id for slot in rng.sample(teacher_slots, min(count, len(teacher_slots))))
            chosen = []
            # پس از انتخاب درصدی از کلاس‌های هر استاد، درصد زبان‌آموزان همان
            # کلاس‌ها نیز جداگانه نمونه‌گیری می‌شود (percentage، مانند حالت قدیمی).
            for slot_id in selected_slot_ids:
                items = enrollments_by_slot[slot_id]
                student_count = max(1, round(len(items) * percentage / 100))
                chosen.extend(rng.sample(items, min(student_count, len(items))))
        else:
            grouped = defaultdict(list)
            for enrollment in enrollments:
                grouped[enrollment.class_slot_id].append(enrollment)
            chosen = []
            for items in grouped.values():
                count = max(1, round(len(items) * percentage / 100))
                chosen.extend(rng.sample(items, min(count, len(items))))
        rows = [_row(item, index) for index, item in enumerate(chosen, start=1)]
        if not rows:
            return Response({'error': 'در این ترم ثبت‌نامی برای نمونه‌گیری وجود ندارد'}, status=400)
        name = str(request.data.get('name') or f'نظرسنجی {term.title} — {percentage}٪').strip()[:200]
        with transaction.atomic():
            survey = ClassSurveyList.objects.create(
                name=name, term=term, percentage=percentage, rows=rows, created_by=request.user,
                use_teacher_class_sampling=use_teacher_class_sampling,
                teacher_class_percentage=teacher_class_percentage if use_teacher_class_sampling else None,
            )
            ClassSurveyResponse.objects.bulk_create([
                ClassSurveyResponse(survey_list=survey, student_id=row['student_id'], class_slot_id=row['class_slot_id'])
                for row in rows
            ])
        return Response({'id': survey.id, 'name': survey.name, 'count': len(rows)}, status=201)

    def delete(self, request):
        if not can_edit_menu(request.user, 'class-management'):
            return Response({'error': 'دسترسی ویرایش ندارید'}, status=403)
        try:
            ClassSurveyList.objects.get(pk=int(request.query_params.get('id'))).delete()
        except (ClassSurveyList.DoesNotExist, ValueError, TypeError):
            return Response({'error': 'لیست پیدا نشد'}, status=404)
        return Response(status=204)


class SurveyResponseView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, survey_id):
        if not _allowed(request):
            return Response({'error': 'دسترسی ندارید'}, status=403)
        try:
            response = ClassSurveyResponse.objects.select_related(
                'student', 'class_slot__term', 'assigned_to', 'survey_list__term'
            ).get(
                survey_list_id=survey_id,
                student_id=int(request.data.get('student_id')),
                class_slot_id=int(request.data.get('class_slot_id')),
            )
        except (ClassSurveyResponse.DoesNotExist, ValueError, TypeError):
            return Response({'error': 'فرد در این لیست پیدا نشد'}, status=404)

        try:
            record_call = _as_bool(request.data.get('record_call'), default=False)
        except ValueError as exc:
            return Response({'error': str(exc)}, status=400)
        call_status = str(request.data.get('call_status') or '').strip()
        if record_call and call_status not in CALL_STATUSES:
            return Response({'error': 'برای ثبت تماس، call_status باید answered یا no_answer باشد'}, status=400)

        dirty = False
        no_answer_call = record_call and call_status == 'no_answer'
        # no_answer هرگز answers موجود را با payload خالی action پاک نمی‌کند.
        if 'answers' in request.data and not no_answer_call:
            if not isinstance(request.data['answers'], dict):
                return Response({'error': 'answers باید یک شیء باشد'}, status=400)
            response.answers = request.data['answers']
            dirty = True
        if 'respondent_name' in request.data:
            respondent_name = str(request.data.get('respondent_name') or '').strip()
            if len(respondent_name) > 200:
                return Response({'error': 'نام پاسخگو نمی‌تواند بیش از ۲۰۰ کاراکتر باشد'}, status=400)
            response.respondent_name = respondent_name
            dirty = True
        if 'assigned_to' in request.data:
            raw_assignee = request.data.get('assigned_to')
            if raw_assignee in (None, '', 0, '0'):
                response.assigned_to = None
            else:
                try:
                    assignee_id = _positive_int(raw_assignee, 'شناسه کارمند')
                except ValueError as exc:
                    return Response({'error': str(exc)}, status=400)
                assignee = User.objects.filter(id=assignee_id, role__in=STAFF_ROLES, is_active=True).first()
                if not assignee:
                    return Response({'error': 'کارمند انتخاب‌شده فعال نیست یا نقش مجاز ندارد'}, status=400)
                response.assigned_to = assignee
            dirty = True

        if record_call:
            response.call_status = call_status
            response.contacted_at = timezone.now()  # زمان کلاینت هرگز پذیرفته نمی‌شود.
            response.completed = False if call_status == 'no_answer' else _response_answers_complete(response)
            dirty = True
        elif 'answers' in request.data and response.call_status == 'answered':
            # autosave answers می‌تواند completion را اصلاح کند، ولی timestamp تماس را هرگز تغییر نمی‌دهد.
            response.completed = _response_answers_complete(response)
            dirty = True

        if dirty:
            response.updated_by = request.user
            response.save()
        return Response({'ok': True, **_response_payload(response, now=timezone.now(), include_session_dates=True)})


class SurveyAssignView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, survey_id):
        if not _allowed(request):
            return Response({'error': 'دسترسی ندارید'}, status=403)
        try:
            survey = ClassSurveyList.objects.get(pk=survey_id)
        except ClassSurveyList.DoesNotExist:
            return Response({'error': 'لیست پیدا نشد'}, status=404)
        mode = str(request.data.get('mode') or '').strip()
        if mode not in ('all', 'distribute'):
            return Response({'error': 'mode باید all یا distribute باشد'}, status=400)
        try:
            employee_ids, staff_by_id = _eligible_staff_by_ids(request.data.get('employee_ids'))
            only_unassigned = _as_bool(request.data.get('only_unassigned'), default=False)
            source_assignee = request.data.get('assigned_to')
            if source_assignee in (None, ''):
                source_assignee_id = None
            else:
                source_assignee_id = _positive_int(source_assignee, 'شناسه کارمند فیلتر')
                if source_assignee_id not in {user.id for user in _staff()}:
                    raise ValueError('کارمند فیلتر فعال نیست یا نقش مجاز ندارد')
        except ValueError as exc:
            return Response({'error': str(exc)}, status=400)
        if mode == 'all' and len(employee_ids) != 1:
            return Response({'error': 'در حالت all دقیقاً یک کارمند انتخاب کنید'}, status=400)

        with transaction.atomic():
            responses = list(
                ClassSurveyResponse.objects.select_for_update().select_related('survey_list')
                .filter(survey_list=survey, completed=False)
            )
            if only_unassigned:
                responses = [response for response in responses if response.assigned_to_id is None]
            if source_assignee_id is not None:
                responses = [response for response in responses if response.assigned_to_id == source_assignee_id]
            row_order = {
                item.get('key'): index
                for index, item in enumerate(survey.rows or [], start=1)
            }
            responses.sort(key=lambda response: row_order.get(f'{response.student_id}:{response.class_slot_id}', 10 ** 12 + response.id))
            counts = {employee_id: 0 for employee_id in employee_ids}
            for index, response in enumerate(responses):
                employee_id = employee_ids[0] if mode == 'all' else employee_ids[index % len(employee_ids)]
                response.assigned_to_id = employee_id
                response.updated_by = request.user
                counts[employee_id] += 1
            if responses:
                ClassSurveyResponse.objects.bulk_update(responses, ['assigned_to', 'updated_by'])

        return Response({
            'ok': True,
            'assigned_count': len(responses),
            'distribution': [
                {'id': employee_id, 'name': _name(staff_by_id[employee_id]), 'count': counts[employee_id]}
                for employee_id in employee_ids
            ],
        })


class SurveyDeadlineView(APIView):
    permission_classes = [IsAuthenticated]

    def patch(self, request, survey_id):
        if not _allowed(request):
            return Response({'error': 'دسترسی ندارید'}, status=403)
        try:
            survey = ClassSurveyList.objects.select_related('term').get(pk=survey_id)
        except ClassSurveyList.DoesNotExist:
            return Response({'error': 'لیست پیدا نشد'}, status=404)
        try:
            enabled = _as_bool(request.data.get('deadline_enabled'), default=False)
        except ValueError as exc:
            return Response({'error': str(exc)}, status=400)

        mode = str(request.data.get('deadline_mode') or 'none').strip()
        if mode not in DEADLINE_MODES:
            return Response({'error': 'deadline_mode نامعتبر است'}, status=400)
        if not enabled or mode == 'none':
            survey.deadline_enabled = False
            survey.deadline_mode = 'none'
            survey.deadline_at = None
            survey.deadline_days = None
            survey.deadline_session_number = None
        elif mode == 'date':
            try:
                survey.deadline_at = _parse_deadline_datetime(request.data.get('deadline_at'))
            except ValueError as exc:
                return Response({'error': str(exc)}, status=400)
            survey.deadline_enabled = True
            survey.deadline_mode = mode
            survey.deadline_days = None
            survey.deadline_session_number = None
        elif mode == 'automatic':
            try:
                days = _positive_int(request.data.get('deadline_days'), 'deadline_days')
            except ValueError as exc:
                return Response({'error': str(exc)}, status=400)
            if days > 365:
                return Response({'error': 'deadline_days باید بین ۱ تا ۳۶۵ باشد'}, status=400)
            survey.deadline_enabled = True
            survey.deadline_mode = mode
            survey.deadline_days = days
            survey.deadline_session_number = None
            # این مقدار برای API و گزارش ذخیره می‌شود؛ محاسبه همچنان از created_at قطعی انجام می‌شود.
            survey.deadline_at = survey.created_at + timedelta(days=days)
        else:  # session
            try:
                session_number = _positive_int(request.data.get('deadline_session_number'), 'deadline_session_number')
            except ValueError as exc:
                return Response({'error': str(exc)}, status=400)
            if session_number > 15:
                return Response({'error': 'deadline_session_number باید بین ۱ تا ۱۵ باشد'}, status=400)
            slots = _survey_slots(survey)
            missing = _missing_session_slots(survey, session_number, slots)
            if missing:
                labels = '، '.join(f'کلاس {slot.number}' for slot in missing[:8])
                suffix = '…' if len(missing) > 8 else ''
                return Response({
                    'error': f'برای جلسه {session_number} تاریخ قابل‌محاسبه نیست؛ {labels}{suffix} تقویم معتبر تا این جلسه ندارند. تاریخ حدس زده نشد.',
                    'session_options': _session_options(survey, slots),
                }, status=400)
            survey.deadline_enabled = True
            survey.deadline_mode = mode
            survey.deadline_at = None
            survey.deadline_days = None
            survey.deadline_session_number = session_number

        with transaction.atomic():
            survey.save(update_fields=[
                'deadline_enabled', 'deadline_mode', 'deadline_at', 'deadline_days', 'deadline_session_number',
            ])
            responses = _save_due_dates(survey)
        return Response({
            'ok': True,
            **_deadline_payload(survey),
            'updated_count': len(responses),
        })


class SurveyDeadlinesView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not _allowed(request):
            return Response({'error': 'دسترسی ندارید'}, status=403)
        # این endpoint برای انتخاب مهلت است؛ نظرسنجی‌های بدون مهلت نیز باید قابل انتخاب باشند.
        payload, error = _list_payload(request, only_deadline_lists=False)
        if error:
            return Response({'error': error}, status=400)
        return Response(payload)


class SurveyDeadlineAlertsView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not _allowed(request):
            return Response({'error': 'دسترسی ندارید'}, status=403)
        now = timezone.now()
        responses = ClassSurveyResponse.objects.filter(completed=False, due_at__isnull=False).select_related(
            'survey_list__term', 'class_slot__term', 'assigned_to'
        )
        # کارمند فقط queue خودش را می‌بیند؛ admin/office و evaluator مجاز، همه را می‌بینند.
        if request.user.role == 'employee':
            responses = responses.filter(assigned_to=request.user)

        grouped = {}
        next_due_at = None
        for response in responses:
            due_at = _effective_due_at(response)
            if not due_at:
                continue
            if due_at >= now and (next_due_at is None or due_at < next_due_at):
                next_due_at = due_at
            assignee = response.assigned_to
            assignee_key = assignee.id if assignee else 'unassigned'
            key = f'{response.survey_list_id}:{assignee_key}:{due_at.isoformat()}'
            if key not in grouped:
                grouped[key] = {
                    'key': key,
                    'survey_id': response.survey_list_id,
                    'survey_name': response.survey_list.name,
                    'assigned_to': assignee.id if assignee else None,
                    'assigned_to_name': _name(assignee),
                    'due_at': _iso(due_at),
                    'due_at_jalali': _jalali_datetime(due_at),
                    'pending_count': 0,
                    'overdue': due_at < now,
                }
            grouped[key]['pending_count'] += 1
        alerts = sorted(grouped.values(), key=lambda item: item['due_at'])
        return Response({
            'server_now': _iso(now),
            'alerts': alerts,
            'next_due_at': _iso(next_due_at),
        })



def _survey_report_responses(request, survey):
    responses = ClassSurveyResponse.objects.filter(
        survey_list=survey, completed=True,
    ).select_related('student', 'class_slot')
    teachers = request.query_params.getlist('teacher')
    levels = request.query_params.getlist('level')
    genders = request.query_params.getlist('gender')
    if teachers:
        teacher_names = [name for name in teachers if name != 'بدون استاد']
        teacher_query = Q(class_slot__teacher_name__in=teacher_names)
        if 'بدون استاد' in teachers:
            teacher_query |= Q(class_slot__teacher_name__isnull=True) | Q(class_slot__teacher_name='')
        responses = responses.filter(teacher_query)
    if levels:
        level_names = [name for name in levels if name != 'بدون سطح']
        level_query = Q(class_slot__assigned_level__in=level_names)
        if 'بدون سطح' in levels:
            level_query |= Q(class_slot__assigned_level__isnull=True) | Q(class_slot__assigned_level='')
        responses = responses.filter(level_query)
    if genders:
        responses = responses.filter(student__gender__in=genders)
    assigned_to = request.query_params.get('assigned_to')
    if request.user.role == 'employee':
        responses = responses.filter(assigned_to=request.user)
    elif assigned_to not in (None, ''):
        try:
            responses = responses.filter(assigned_to_id=int(assigned_to))
        except (TypeError, ValueError):
            return None, Response({'error': 'شناسه کارمند نامعتبر است'}, status=400)
    return responses, None


def _survey_report_filter_options(survey):
    levels_by_teacher = defaultdict(set)
    for teacher_name, level_name in ClassSlot.objects.filter(term_id=survey.term_id).values_list(
        'teacher_name', 'assigned_level',
    ).distinct():
        teacher_label = (teacher_name or '').strip() or 'بدون استاد'
        level_label = (level_name or '').strip() or 'بدون سطح'
        levels_by_teacher[teacher_label].add(level_label)
    normalized = {
        teacher: sorted(levels, key=lambda value: value.casefold())
        for teacher, levels in levels_by_teacher.items()
    }
    return {
        'teacher_options': sorted(normalized, key=lambda value: value.casefold()),
        'level_options': sorted({level for levels in normalized.values() for level in levels}, key=lambda value: value.casefold()),
        'levels_by_teacher': normalized,
    }


def _build_survey_report(response_items):
    groups = defaultdict(lambda: {'sum': 0, 'count': 0, 'yes_count': 0, 'no_count': 0, 'teacher': '', 'level': ''})
    teacher_groups = defaultdict(lambda: {
        'teacher': '', 'response_count': 0, 'yes_count': 0, 'no_count': 0, 'count': 0,
        'age_groups': defaultdict(lambda: {'response_count': 0, 'yes_count': 0, 'no_count': 0, 'count': 0}),
    })
    question_groups = defaultdict(lambda: {
        'yes_count': 0, 'no_count': 0, 'count': 0, 'teacher': '', 'level': '',
        'age_group': '', 'question': '', 'question_id': None,
    })
    class_age_groups = defaultdict(lambda: {
        'teacher': '', 'level': '', 'age_group': '', 'response_count': 0,
        'yes_count': 0, 'no_count': 0, 'count': 0,
    })
    question_map = {str(question.id): question for question in ClassSurveyQuestion.objects.all()}

    for item in response_items:
        slot = item.class_slot
        teacher_name = (slot.teacher_name or '').strip() or 'بدون استاد'
        level_name = (slot.assigned_level or '').strip() or 'بدون سطح'
        age_group = _age_group(slot.assigned_level or item.student.language_level or '')
        class_key = f'{teacher_name} — {level_name}'
        groups[class_key].update({'teacher': teacher_name, 'level': level_name})
        teacher = teacher_groups[teacher_name]
        teacher['teacher'] = teacher_name
        teacher['response_count'] += 1
        if age_group in ('kids', 'teen', 'adult'):
            teacher['age_groups'][age_group]['response_count'] += 1
            class_age_key = (teacher_name, level_name, age_group)
            class_age_groups[class_age_key].update({
                'teacher': teacher_name, 'level': level_name, 'age_group': age_group,
            })
            class_age_groups[class_age_key]['response_count'] += 1

        for raw_question_id, raw_value in (item.answers or {}).items():
            answer = raw_value.get('answer') if isinstance(raw_value, dict) else raw_value
            if answer not in ('yes', 'no'):
                continue
            question_id = str(raw_question_id)
            question = question_map.get(question_id)
            question_text = question.text if question else f'سؤال {question_id}'
            question_key = (teacher_name, level_name, age_group, question_id)
            question_row = question_groups[question_key]
            question_row.update({
                'teacher': teacher_name, 'level': level_name, 'age_group': age_group,
                'question': question_text, 'question_id': question.id if question else question_id,
            })
            count_field = 'yes_count' if answer == 'yes' else 'no_count'
            question_row[count_field] += 1
            question_row['count'] += 1
            groups[class_key][count_field] += 1
            groups[class_key]['sum'] += 1 if answer == 'yes' else -1
            groups[class_key]['count'] += 1
            teacher[count_field] += 1
            teacher['count'] += 1
            if age_group in ('kids', 'teen', 'adult'):
                teacher['age_groups'][age_group][count_field] += 1
                teacher['age_groups'][age_group]['count'] += 1
                class_age_groups[class_age_key][count_field] += 1
                class_age_groups[class_age_key]['count'] += 1

    result = []
    for key, value in groups.items():
        result.append({
            'key': key, **value,
            'score': round(value['sum'] / value['count'] * 100, 1) if value['count'] else 0,
        })
    result.sort(key=lambda item: item['score'], reverse=True)

    teacher_summary = []
    for teacher_name, values in teacher_groups.items():
        total_answers = values['yes_count'] + values['no_count']
        age_groups = {}
        for age_key in ('kids', 'teen', 'adult'):
            age_values = values['age_groups'][age_key]
            age_count = age_values['yes_count'] + age_values['no_count']
            age_groups[age_key] = {
                **age_values,
                'satisfaction_percent': round(age_values['yes_count'] * 100 / age_count, 1) if age_count else 0,
            }
        teacher_summary.append({
            'key': teacher_name, 'teacher': teacher_name,
            'response_count': values['response_count'], 'yes_count': values['yes_count'],
            'no_count': values['no_count'], 'count': values['count'],
            'satisfaction_percent': round(values['yes_count'] * 100 / total_answers, 1) if total_answers else 0,
            'age_groups': age_groups,
        })
    teacher_summary.sort(key=lambda item: (-item['satisfaction_percent'], item['teacher']))

    question_summary = []
    for key, values in question_groups.items():
        total = values['count']
        question_summary.append({
            'key': '|'.join(map(str, key)), **values,
            'yes_percent': round(values['yes_count'] * 100 / total, 1) if total else 0,
            'no_percent': round(values['no_count'] * 100 / total, 1) if total else 0,
            'satisfaction_percent': round(values['yes_count'] * 100 / total, 1) if total else 0,
        })
    question_summary.sort(key=lambda item: (item['teacher'], item['level'], item['age_group'], str(item['question'])))

    class_age_summary = []
    for key, values in class_age_groups.items():
        total = values['yes_count'] + values['no_count']
        class_age_summary.append({
            'key': '|'.join(key), **values,
            'yes_percent': round(values['yes_count'] * 100 / total, 1) if total else 0,
            'no_percent': round(values['no_count'] * 100 / total, 1) if total else 0,
            'satisfaction_percent': round(values['yes_count'] * 100 / total, 1) if total else 0,
        })
    class_age_summary.sort(key=lambda item: (item['teacher'], item['level'], item['age_group']))

    yes_count = sum(item['yes_count'] for item in teacher_summary)
    no_count = sum(item['no_count'] for item in teacher_summary)
    total_answers = yes_count + no_count
    overall_age_groups = {}
    for age_key in ('kids', 'teen', 'adult'):
        values = {
            field: sum(item['age_groups'][age_key][field] for item in teacher_summary)
            for field in ('response_count', 'yes_count', 'no_count', 'count')
        }
        age_count = values['yes_count'] + values['no_count']
        overall_age_groups[age_key] = {
            **values,
            'satisfaction_percent': round(values['yes_count'] * 100 / age_count, 1) if age_count else 0,
        }
    overall_summary = {
        'key': 'overall', 'teacher': 'مجموع همه استادان',
        'response_count': sum(item['response_count'] for item in teacher_summary),
        'yes_count': yes_count, 'no_count': no_count, 'count': total_answers,
        'satisfaction_percent': round(yes_count * 100 / total_answers, 1) if total_answers else 0,
        'age_groups': overall_age_groups,
    }
    return {
        'groups': result, 'teacher_summary': teacher_summary,
        'overall_summary': overall_summary, 'question_summary': question_summary,
        'class_age_summary': class_age_summary,
        'total_completed': len(response_items),
    }


class SurveyReportView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, survey_id):
        if not _allowed(request):
            return Response({'error': 'دسترسی ندارید'}, status=403)
        survey = ClassSurveyList.objects.filter(pk=survey_id).first()
        if not survey:
            return Response({'error': 'لیست نظرسنجی پیدا نشد'}, status=404)
        responses, error = _survey_report_responses(request, survey)
        if error:
            return error
        payload = _build_survey_report(list(responses))
        payload.update(_survey_report_filter_options(survey))
        return Response(payload)


class SurveyReportExcelView(APIView):
    """فایل XLSX فیلترپذیرِ گزارش استاد، سطح، گروه سنی و پاسخ هر سؤال."""
    permission_classes = [IsAuthenticated]

    def get(self, request, survey_id):
        if not _allowed(request):
            return Response({'error': 'دسترسی ندارید'}, status=403)
        survey = ClassSurveyList.objects.select_related('term').filter(pk=survey_id).first()
        if not survey:
            return Response({'error': 'لیست نظرسنجی پیدا نشد'}, status=404)
        responses, error = _survey_report_responses(request, survey)
        if error:
            return error
        response_items = list(responses)
        report = _build_survey_report(response_items)
        try:
            from io import BytesIO
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            return Response({'error': 'کتابخانهٔ openpyxl روی سرور نصب نیست؛ برای ساخت خروجی XLSX آن را نصب کنید.'}, status=503)

        workbook = Workbook()
        summary = workbook.active
        summary.title = 'رضایت استادان'
        summary.append([
            'ترم', 'استاد', 'نظرسنجی تکمیل‌شده', 'پاسخ بله', 'پاسخ خیر', 'مجموع پاسخ‌ها',
            'درصد رضایت کل', 'رضایت کودک ٪', 'رضایت نوجوان ٪', 'رضایت بزرگسال ٪',
        ])
        summary_rows = [*report['teacher_summary'], report['overall_summary']]
        for item in summary_rows:
            ages = item.get('age_groups', {})
            summary.append([
                survey.term.title, item['teacher'], item['response_count'], item['yes_count'],
                item['no_count'], item['count'], item['satisfaction_percent'],
                ages.get('kids', {}).get('satisfaction_percent', 0),
                ages.get('teen', {}).get('satisfaction_percent', 0),
                ages.get('adult', {}).get('satisfaction_percent', 0),
            ])

        questions = workbook.create_sheet('نتایج هر سؤال')
        questions.append([
            'ترم', 'استاد', 'سطح', 'گروه سنی', 'سؤال', 'تعداد بله', 'درصد بله',
            'تعداد خیر', 'درصد خیر', 'مجموع پاسخ‌ها', 'درصد رضایت',
        ])
        age_labels = {'kids': 'کودک', 'teen': 'نوجوان', 'adult': 'بزرگسال', 'all': 'نامشخص'}
        for item in report['question_summary']:
            questions.append([
                survey.term.title, item['teacher'], item['level'], age_labels.get(item['age_group'], item['age_group']),
                item['question'], item['yes_count'], item['yes_percent'], item['no_count'],
                item['no_percent'], item['count'], item['satisfaction_percent'],
            ])

        class_age = workbook.create_sheet('رضایت کلاس و رده سنی')
        class_age.append([
            'ترم', 'استاد', 'سطح کلاس', 'گروه سنی', 'نظرسنجی تکمیل‌شده',
            'تعداد بله', 'درصد بله', 'تعداد خیر', 'درصد خیر', 'مجموع پاسخ‌ها', 'درصد رضایت',
        ])
        for item in report['class_age_summary']:
            class_age.append([
                survey.term.title, item['teacher'], item['level'], age_labels.get(item['age_group'], item['age_group']),
                item['response_count'], item['yes_count'], item['yes_percent'],
                item['no_count'], item['no_percent'],
                item['count'], item['satisfaction_percent'],
            ])

        detail = workbook.create_sheet('جزئیات پاسخ‌ها')
        detail.append([
            'ترم', 'استاد', 'سطح', 'گروه سنی', 'نام زبان‌آموز', 'شماره تماس',
            'پاسخ بله', 'پاسخ خیر', 'مجموع پاسخ‌ها', 'درصد رضایت',
        ])
        for item in response_items:
            yes_count = no_count = 0
            for value in (item.answers or {}).values():
                answer = value.get('answer') if isinstance(value, dict) else value
                if answer == 'yes': yes_count += 1
                elif answer == 'no': no_count += 1
            count = yes_count + no_count
            age = _age_group(item.class_slot.assigned_level or item.student.language_level or '')
            detail.append([
                survey.term.title, (item.class_slot.teacher_name or '').strip() or 'بدون استاد',
                (item.class_slot.assigned_level or '').strip() or 'بدون سطح', age_labels.get(age, age),
                f'{item.student.first_name} {item.student.last_name}'.strip(), item.student.phone or '',
                yes_count, no_count, count, round(yes_count * 100 / count, 1) if count else 0,
            ])

        for sheet in (summary, questions, class_age, detail):
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = f'A1:{sheet.cell(row=1, column=sheet.max_column).column_letter}{sheet.max_row}'
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='17365D')
                cell.alignment = Alignment(horizontal='right')
            for column in sheet.columns:
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = min(48, max(14, max(len(str(cell.value or '')) for cell in column) + 2))
            sheet.sheet_view.rightToLeft = True
        if summary.max_row > 1:
            for cell in summary[summary.max_row]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill('solid', fgColor='DCE6F1')

        buffer = BytesIO()
        workbook.save(buffer)
        response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="class-survey-teachers-{survey.term_id}.xlsx"'
        return response

class SurveyQuestionView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not _allowed(request):
            return Response({'error': 'دسترسی ندارید'}, status=403)
        return Response(_questions(request.query_params.get('age_group')))

    def post(self, request):
        if not can_edit_menu(request.user, 'class-management'):
            return Response({'error': 'دسترسی ویرایش ندارید'}, status=403)
        text = str(request.data.get('text') or '').strip()
        if not text:
            return Response({'error': 'متن سوال الزامی است'}, status=400)
        question = ClassSurveyQuestion.objects.create(
            age_group=request.data.get('age_group') or 'all', text=text,
            sort_order=int(request.data.get('sort_order') or 0),
        )
        return Response({
            'id': question.id, 'age_group': question.age_group, 'text': question.text,
            'sort_order': question.sort_order,
        }, status=201)

    def patch(self, request, pk):
        if not can_edit_menu(request.user, 'class-management'):
            return Response({'error': 'دسترسی ویرایش ندارید'}, status=403)
        question = ClassSurveyQuestion.objects.get(pk=pk)
        question.text = str(request.data.get('text', question.text)).strip()
        question.age_group = request.data.get('age_group', question.age_group)
        question.active = bool(request.data.get('active', question.active))
        question.save()
        return Response({'ok': True})

    def delete(self, request, pk):
        if not can_edit_menu(request.user, 'class-management'):
            return Response({'error': 'دسترسی ویرایش ندارید'}, status=403)
        ClassSurveyQuestion.objects.filter(pk=pk).delete()
        return Response(status=204)
